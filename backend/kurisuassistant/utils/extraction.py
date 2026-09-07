"""Turn a drive file into text the retrieval index can chunk (#6).

Plain text, Markdown, source code and JSON need nothing but a UTF-8 decode. The
office formats each need a library, imported lazily inside the function that
uses it — the same rule every heavy provider follows, so the collection-time
import graph (and ``requirements-ci.txt``) stays small.

What comes back is a list of ``Page`` objects: the text, and the page it came
from where the format has pages (PDF pages, PPTX slides, XLSX sheets) or ``None``
where it does not. A citation then says "page 3" for a PDF and "lines 40–58"
for a text file.

Whether a file is text is decided by looking at the bytes, not only the name:
``mimetypes`` calls ``.ts`` a video and ``.toml`` an octet-stream, and a file
with no extension is often a README. The sniff is the one ``drive_storage``
uses to refuse binaries in ``drive_read`` — no NUL byte, decodes as UTF-8.

**Secrets are never indexed.** A dotenv file, a private key or a certificate
bundle is text and would sniff as such, but a retrieval index is a second copy
of its contents that every recall query can surface. Names matching
``SECRET_NAMES`` and ``SECRET_EXTENSIONS`` are refused before the bytes are read.

Two failure kinds, both deterministic and therefore not retried by the indexer:
``Unextractable`` (a binary, an unknown format, a secret) and ``ExtractionError``
(a file that claims a format and cannot be parsed as it).
"""

from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

logger = logging.getLogger(__name__)

#: How many bytes are sniffed to decide whether an unknown file is text.
SNIFF_BYTES = 8192

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".adoc", ".org", ".tex", ".csv", ".tsv",
    ".log", ".ini", ".cfg", ".conf", ".toml", ".yaml", ".yml", ".json", ".jsonl",
    ".xml", ".properties", ".srt", ".vtt",
    ".py", ".pyi", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".kt", ".kts",
    ".java", ".scala", ".go", ".rs", ".c", ".h", ".cc", ".cpp", ".hpp", ".cs",
    ".swift", ".m", ".rb", ".php", ".pl", ".lua", ".sh", ".bash", ".zsh", ".fish",
    ".ps1", ".bat", ".sql", ".r", ".jl", ".dart", ".vue", ".svelte", ".css",
    ".scss", ".less", ".gradle", ".cmake", ".mk", ".dockerfile", ".proto",
    ".graphql", ".gql", ".tf", ".hcl", ".nix",
}
TEXT_MIMES = {
    "application/json", "application/xml", "application/x-yaml", "application/yaml",
    "application/javascript", "application/x-sh", "application/toml",
    "application/x-httpd-php", "application/sql",
}
HTML_EXTENSIONS = {".html", ".htm", ".xhtml"}
PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
PPTX_EXTENSIONS = {".pptx"}
XLSX_EXTENSIONS = {".xlsx", ".xlsm"}

#: Files that are text and must still never be indexed. Matched on the lowercased
#: file name (prefix for the dotenv family, exact for the rest).
SECRET_NAME_PREFIXES = (".env",)
SECRET_NAMES = {"id_rsa", "id_ed25519", "id_ecdsa", "id_dsa", ".netrc", ".npmrc", ".pypirc",
                "credentials", ".htpasswd", "secrets.yaml", "secrets.yml", "secrets.json"}
SECRET_EXTENSIONS = {".pem", ".key", ".p12", ".pfx", ".jks", ".keystore", ".kdbx", ".asc", ".gpg"}


class ExtractionError(Exception):
    """The file names a format this module knows and cannot be read as it."""


class Unextractable(ExtractionError):
    """Nothing here will read this file: a binary, an unknown format, a secret."""


@dataclass(frozen=True)
class Page:
    text: str
    page: Optional[int]


def looks_like_text(sample: bytes) -> bool:
    """The check ``drive_read`` applies: no NUL byte, and it decodes as UTF-8.

    A sample cut mid-character is still text; only a decode error before the
    last few bytes says otherwise.
    """
    if not sample:
        return True
    if b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
        return True
    except UnicodeDecodeError as e:
        return e.start >= len(sample) - 4


def is_secret_name(name: str) -> bool:
    lowered = name.lower()
    if lowered.startswith(SECRET_NAME_PREFIXES) or lowered in SECRET_NAMES:
        return True
    return os.path.splitext(lowered)[1] in SECRET_EXTENSIONS


def kind_of(name: str, mime: Optional[str]) -> Optional[str]:
    """Which reader a file gets from its name and guessed MIME type.

    ``None`` means "not known from the name" — the caller then sniffs the bytes
    and reads it as text if it is text. So this never says "no": it only says
    what it can tell without opening the file.
    """
    ext = os.path.splitext(name)[1].lower()
    if ext in PDF_EXTENSIONS:
        return "pdf"
    if ext in DOCX_EXTENSIONS:
        return "docx"
    if ext in PPTX_EXTENSIONS:
        return "pptx"
    if ext in XLSX_EXTENSIONS:
        return "xlsx"
    if ext in HTML_EXTENSIONS or mime == "text/html":
        return "html"
    if ext in TEXT_EXTENSIONS or (mime and (mime.startswith("text/") or mime in TEXT_MIMES)):
        return "text"
    return None


def _decode(data: bytes) -> str:
    if not looks_like_text(data):
        raise Unextractable("binary")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as e:
        raise Unextractable("not UTF-8") from e
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _read_capped(path: Path, max_bytes: int) -> bytes:
    with open(path, "rb") as fh:
        data = fh.read(max_bytes + 1)
    if len(data) > max_bytes:
        raise Unextractable(f"larger than {max_bytes} bytes")
    return data


def _extract_text(path: Path, max_bytes: int) -> List[Page]:
    text = _decode(_read_capped(path, max_bytes))
    return [Page(text, None)] if text.strip() else []


def _extract_html(path: Path, max_bytes: int) -> List[Page]:
    from html.parser import HTMLParser

    raw = _decode(_read_capped(path, max_bytes))

    block_tags = {
        "p", "div", "br", "li", "ul", "ol", "h1", "h2", "h3", "h4", "h5", "h6",
        "tr", "table", "section", "article", "header", "footer", "pre", "blockquote",
        "hr", "dd", "dt", "figcaption", "title",
    }

    class _Text(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.parts: List[str] = []
            self._skip = 0

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style", "noscript"):
                self._skip += 1
            elif tag in block_tags:
                self.parts.append("\n")

        def handle_endtag(self, tag):
            if tag in ("script", "style", "noscript"):
                self._skip = max(0, self._skip - 1)
            elif tag in block_tags:
                self.parts.append("\n")

        def handle_data(self, data):
            if not self._skip:
                self.parts.append(data)

    parser = _Text()
    parser.feed(raw)
    parser.close()
    lines = [" ".join(line.split()) for line in "".join(parser.parts).split("\n")]
    text = "\n".join(line for line in lines if line)
    return [Page(text, None)] if text.strip() else []


def _extract_pdf(path: Path) -> List[Page]:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as e:  # pragma: no cover - dependency missing
        raise ExtractionError("pypdf is not installed") from e

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            # An owner password only; a real user password cannot be guessed.
            if reader.decrypt("") == 0:
                raise ExtractionError("encrypted PDF")
        pages = []
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(Page(text, number))
        return pages
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as e:
        raise ExtractionError(f"unreadable PDF: {e}") from e


def _extract_docx(path: Path) -> List[Page]:
    try:
        import docx
    except ImportError as e:  # pragma: no cover
        raise ExtractionError("python-docx is not installed") from e

    try:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
    except Exception as e:
        raise ExtractionError(f"unreadable DOCX: {e}") from e
    text = "\n".join(parts)
    return [Page(text, None)] if text.strip() else []


def _extract_pptx(path: Path) -> List[Page]:
    try:
        from pptx import Presentation
    except ImportError as e:  # pragma: no cover
        raise ExtractionError("python-pptx is not installed") from e

    try:
        presentation = Presentation(str(path))
        pages = []
        for number, slide in enumerate(presentation.slides, start=1):
            parts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append("\t".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
            if text.strip():
                pages.append(Page(text, number))
        return pages
    except Exception as e:
        raise ExtractionError(f"unreadable PPTX: {e}") from e


def _extract_xlsx(path: Path) -> List[Page]:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise ExtractionError("openpyxl is not installed") from e

    try:
        # A file object, not the path: the blob has no extension and openpyxl
        # refuses a path whose name does not end in .xlsx. Read-only mode reads
        # lazily, so the file stays open until the last sheet is walked.
        pages = []
        with open(path, "rb") as fh:
            workbook = openpyxl.load_workbook(fh, read_only=True, data_only=True)
            try:
                for number, sheet in enumerate(workbook.worksheets, start=1):
                    rows = [f"Sheet: {sheet.title}"]
                    for row in sheet.iter_rows(values_only=True):
                        cells = ["" if v is None else str(v) for v in row]
                        if any(cells):
                            rows.append("\t".join(cells).rstrip())
                    if len(rows) > 1:
                        pages.append(Page("\n".join(rows), number))
            finally:
                workbook.close()
        return pages
    except Exception as e:
        raise ExtractionError(f"unreadable XLSX: {e}") from e


def extract(path: Path, name: str, mime: Optional[str], max_bytes: int) -> List[Page]:
    """Text for a drive file, page by page where the format has pages.

    Raises ``Unextractable`` for a binary, an unknown format or a secret, and
    ``ExtractionError`` for a known format that cannot be parsed. Files over
    ``max_bytes`` are refused before they are read; the office readers are given
    the whole file because their libraries need the container intact.
    """
    path = Path(path)
    if is_secret_name(name):
        raise Unextractable("looks like a secret; not indexed")
    if path.stat().st_size > max_bytes:
        raise Unextractable(f"larger than {max_bytes} bytes")

    kind = kind_of(name, mime)
    if kind is None:
        with open(path, "rb") as fh:
            sample = fh.read(SNIFF_BYTES)
        kind = "text" if looks_like_text(sample) else None
    if kind is None:
        raise Unextractable("binary")

    if kind == "text":
        return _extract_text(path, max_bytes)
    if kind == "html":
        return _extract_html(path, max_bytes)
    if kind == "pdf":
        return _extract_pdf(path)
    if kind == "docx":
        return _extract_docx(path)
    if kind == "pptx":
        return _extract_pptx(path)
    if kind == "xlsx":
        return _extract_xlsx(path)
    raise Unextractable(kind)  # pragma: no cover - every kind above is handled
